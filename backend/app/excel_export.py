"""生成 .xlsx 导出文件(openpyxl 三工作表:测试概述 / 功能点清单 / 测试用例)。

对齐 ai-testgen skill 报告结构,批量查询避免 N+1(与 xmind_export 同风格)。
"""

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy.orm import Session, selectinload

from app.models import Case, FeaturePoint, Module, Project

# 执行结果映射
_EXEC_RESULT = {True: "通过", False: "未通过", None: "未执行"}

# 样式常量
_HEADER_FILL = PatternFill(start_color="409EFF", end_color="409EFF", fill_type="solid")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_DATA_ALIGN = Alignment(wrap_text=True, vertical="top")

# 公式注入危险前缀:值以这些字符开头时可能被 Excel 解释为公式
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def _safe_cell(ws, row: int, column: int, value) -> None:
    """写入单元格,用户可控字符串做公式注入防护:危险前缀强制文本类型。"""
    cell = ws.cell(row=row, column=column, value=value)
    if isinstance(value, str) and value and value[0] in _FORMULA_PREFIXES:
        cell.data_type = "s"


def _build_module_path(module: Module, module_map: dict[int, Module]) -> str:
    """递归向上拼模块路径,如 父模块/子模块。"""
    parts: list[str] = []
    cur: Module | None = module
    while cur is not None:
        parts.append(cur.name)
        cur = module_map.get(cur.parent_id) if cur.parent_id else None
    parts.reverse()
    return "/".join(parts)


def _format_steps(case: Case) -> str:
    """将步骤合并为多行文本:1. 操作\n预期: xxx"""
    lines: list[str] = []
    for step in case.steps:
        lines.append(f"{step.step_no}. {step.action}\n预期: {step.expected}")
    return "\n".join(lines)


def _write_overview(ws, project: Project, case_count: int, fp_count: int) -> None:
    """写入「测试概述」工作表。"""
    rows = [
        ("项目名", project.name),
        ("导出时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("用例数", case_count),
        ("功能点数", fp_count),
    ]
    for i, (label, value) in enumerate(rows, start=1):
        ws.cell(row=i, column=1, value=label)
        _safe_cell(ws, i, 2, value)
    ws.column_dimensions["A"].width = 15
    ws.column_dimensions["B"].width = 40


def _write_feature_points(ws, fp_data: list[dict]) -> None:
    """写入「功能点清单」工作表。"""
    headers = ["功能点ID", "功能名称", "所属模块路径", "用例数"]
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _DATA_ALIGN
    for r, fp in enumerate(fp_data, start=2):
        for c, val in enumerate([fp["id"], fp["name"], fp["module_path"], fp["case_count"]], start=1):
            _safe_cell(ws, r, c, val)
            ws.cell(row=r, column=c).alignment = _DATA_ALIGN
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 30
    ws.column_dimensions["D"].width = 10


def _write_cases(ws, case_rows: list[dict]) -> None:
    """写入「测试用例」工作表。"""
    headers = ["用例ID", "功能点ID", "标题", "优先级", "前置条件", "步骤", "备注", "执行结果"]
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _DATA_ALIGN
    col_widths = {3: 40, 5: 30, 6: 60, 7: 30}
    for c, w in col_widths.items():
        ws.column_dimensions[chr(64 + c)].width = w  # C=3, E=5, F=6, G=7
    for r, case in enumerate(case_rows, start=2):
        vals = [case["id"], case["fp_id"], case["title"], case["priority"],
                case["precondition"], case["steps"], case["remark"], case["exec_result"]]
        for c, val in enumerate(vals, start=1):
            _safe_cell(ws, r, c, val)
            ws.cell(row=r, column=c).alignment = _DATA_ALIGN
    ws.freeze_panes = "A2"


def build_excel_bytes(db: Session, project: Project) -> bytes:
    """构建三工作表 xlsx 并返回 bytes。"""
    # 三次批量查询(与 xmind_export 同风格)
    modules = (
        db.query(Module)
        .filter(Module.project_id == project.id)
        .order_by(Module.sort_order, Module.id)
        .all()
    )
    module_ids = [m.id for m in modules]

    fps: list[FeaturePoint] = []
    cases: list[Case] = []
    if module_ids:
        fps = (
            db.query(FeaturePoint)
            .filter(FeaturePoint.module_id.in_(module_ids))
            .order_by(FeaturePoint.sort_order, FeaturePoint.id)
            .all()
        )
    fp_ids = [f.id for f in fps]
    if fp_ids:
        cases = (
            db.query(Case)
            .filter(Case.feature_point_id.in_(fp_ids))
            .options(selectinload(Case.steps))
            .order_by(Case.sort_order, Case.id)
            .all()
        )

    # 建立索引
    module_map: dict[int, Module] = {m.id: m for m in modules}
    fps_by_module: dict[int, list[FeaturePoint]] = {}
    for f in fps:
        fps_by_module.setdefault(f.module_id, []).append(f)
    cases_by_fp: dict[int, list[Case]] = {}
    for c in cases:
        cases_by_fp.setdefault(c.feature_point_id, []).append(c)

    # 构建功能点清单数据
    fp_data: list[dict] = []
    for f in fps:
        mod = module_map.get(f.module_id)
        mod_name = mod.name if mod else ""
        fp_data.append({
            "id": f.id,
            "name": f.name,
            "module_path": _build_module_path(mod, module_map) if mod else mod_name,
            "case_count": len(cases_by_fp.get(f.id, [])),
        })

    # 构建用例行数据
    case_rows: list[dict] = []
    for c in cases:
        case_rows.append({
            "id": c.id,
            "fp_id": c.feature_point_id,
            "title": c.title,
            "priority": c.priority,
            "precondition": c.precondition,
            "steps": _format_steps(c),
            "remark": c.remark,
            "exec_result": _EXEC_RESULT.get(c.executed_pass, "未执行"),
        })

    # 写工作簿
    wb = Workbook()
    _write_overview(wb.active, project, len(case_rows), len(fps))
    wb.active.title = "测试概述"

    ws_fp = wb.create_sheet("功能点清单")
    _write_feature_points(ws_fp, fp_data)
    ws_fp.freeze_panes = "A2"

    ws_case = wb.create_sheet("测试用例")
    _write_cases(ws_case, case_rows)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

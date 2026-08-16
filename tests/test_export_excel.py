# tests/test_export_excel.py
import io

from openpyxl import load_workbook

import pytest


@pytest.fixture()
def seeded_pid(client):
    """与 test_export_xmind.py 的 seeded_pid 同构:项目 → 登录模块 → 子模块 + 功能点 → P0 用例(1 步,已执行通过)。"""
    pid = client.post("/api/projects", json={"name": "导出测试项目"}).json()["id"]
    m1 = client.post(f"/api/projects/{pid}/modules", json={"name": "登录模块"}).json()
    client.post(f"/api/projects/{pid}/modules", json={"name": "子模块", "parent_id": m1["id"]})
    fp = client.post(f"/api/modules/{m1['id']}/feature-points", json={"name": "账号登录"}).json()
    case = client.post(
        f"/api/feature-points/{fp['id']}/cases",
        json={
            "title": "正确账号密码登录成功",
            "priority": "P0",
            "precondition": "已注册用户",
            "remark": "冒烟用例",
            "steps": [{"action": "输入账号密码", "expected": "登录成功"}],
        },
    ).json()
    client.patch(f"/api/cases/{case['id']}/execution", json={"executed_pass": True})
    return pid


def test_export_excel_structure(client, seeded_pid):
    resp = client.get(f"/api/projects/{seeded_pid}/export/excel")
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "attachment" in resp.headers["content-disposition"]

    wb = load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == ["测试概述", "功能点清单", "测试用例"]

    overview = wb["测试概述"]
    cells = {overview.cell(row=r, column=1).value: overview.cell(row=r, column=2).value for r in range(1, overview.max_row + 1)}
    assert cells["项目名"] == "导出测试项目"
    assert cells["功能点数"] == 1
    assert cells["用例数"] == 1

    fps = wb["功能点清单"]
    assert [fps.cell(row=1, column=c).value for c in range(1, 5)] == ["功能点ID", "功能名称", "所属模块路径", "用例数"]
    assert fps.cell(row=2, column=2).value == "账号登录"
    assert fps.cell(row=2, column=3).value == "登录模块"
    assert fps.cell(row=2, column=4).value == 1

    ws = wb["测试用例"]
    assert ws.max_row >= 2
    assert [ws.cell(row=1, column=c).value for c in range(1, 9)] == [
        "用例ID", "功能点ID", "标题", "优先级", "前置条件", "步骤", "备注", "执行结果",
    ]
    assert ws.cell(row=2, column=3).value == "正确账号密码登录成功"
    assert ws.cell(row=2, column=4).value == "P0"
    assert ws.cell(row=2, column=5).value == "已注册用户"
    assert "1. 输入账号密码" in ws.cell(row=2, column=6).value
    assert "预期: 登录成功" in ws.cell(row=2, column=6).value
    assert ws.cell(row=2, column=7).value == "冒烟用例"
    assert ws.cell(row=2, column=8).value == "通过"


def test_export_excel_project_not_found(client):
    assert client.get("/api/projects/99999/export/excel").status_code == 404


def test_export_excel_empty_project(client):
    pid = client.post("/api/projects", json={"name": "空导出项目"}).json()["id"]
    resp = client.get(f"/api/projects/{pid}/export/excel")
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == ["测试概述", "功能点清单", "测试用例"]
    overview = wb["测试概述"]
    cells = {overview.cell(row=r, column=1).value: overview.cell(row=r, column=2).value for r in range(1, overview.max_row + 1)}
    assert cells["项目名"] == "空导出项目"
    assert cells["用例数"] == 0
    assert wb["功能点清单"].max_row == 1  # 仅表头
    assert wb["测试用例"].max_row == 1
